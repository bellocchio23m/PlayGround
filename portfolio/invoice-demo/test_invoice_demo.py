import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
import invoice_demo as demo


class InvoiceChecks(unittest.TestCase):
    def example(self, name='example-en.pdf'):
        layout, lines = demo.EXAMPLES[name]
        return demo.parse_text('\n'.join(lines), layout)

    def test_actual_pdf_extraction_matches_source(self):
        record = demo.parse_pdf(demo.ROOT / 'example-en.pdf', 'en')
        self.assertEqual(record['invoice'], '0007')
        self.assertEqual(record['items'][0]['unit_raw'], '12.50')
        self.assertEqual(len(record['items']), 2)
        self.assertEqual(record['issues'], [])

    def test_decimal_comma_and_bad_source_total_preserved(self):
        record = demo.parse_pdf(demo.ROOT / 'example-it.pdf', 'it')
        self.assertEqual(record['total_raw'], '43,00')
        self.assertEqual(record['items'][0]['gross_raw'], '36,60')
        self.assertEqual(record['issues'][0]['sum_of_lines'], '42.70')
        self.assertEqual(record['issues'][0]['type'], 'invoice_total_mismatch')

    def test_missing_header_is_not_invented(self):
        text = '\n'.join(demo.EXAMPLES['example-en.pdf'][1]).replace('Invoice No: 0007', '')
        with self.assertRaises(ValueError):
            demo.parse_text(text, 'en')

    def test_ambiguous_number_rejected(self):
        with self.assertRaises(ValueError):
            demo.amount('1,234.56', '.')

    def test_workbook_roundtrip_preserves_ids_and_literal_strings(self):
        record = demo.parse_pdf(demo.ROOT/'example-en.pdf', 'en')
        record['supplier'] = '=1+1'
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'out.xlsx'
            demo.write_workbook([record, record], path)
            wb = load_workbook(path)
            self.assertEqual(wb['0007']['B3'].value, '0007')
            self.assertEqual(wb['0007']['B5'].value, '=1+1')
            self.assertEqual(wb['0007']['B5'].data_type, 's')
            self.assertIn('0007_2', wb.sheetnames)

    def test_workbook_flags_source_mismatch(self):
        wb = load_workbook(demo.ROOT/'invoice-review.xlsx')
        self.assertEqual(wb['0008']['B7'].value, '43,00')
        self.assertEqual(wb['0008']['B7'].fill.fgColor.rgb, '00FFF2CC')
        self.assertEqual(wb['Review']['D4'].value, 'REVIEW REQUIRED')


if __name__ == '__main__':
    unittest.main(verbosity=2)
