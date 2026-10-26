"""Two fictional, agreed one-page layouts; not a general invoice parser.

Run: python invoice_demo.py
Dependencies: reportlab, pdfplumber, openpyxl. No network requests.
Original demonstration by Marco Bellocchio, not commissioned work.
"""
import hashlib
import json
import re
from decimal import Decimal
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.pdfgen import canvas

ROOT = Path(__file__).parent
LAYOUTS = {
    'en': ('Invoice No:', 'Date:', 'Supplier:', 'Currency:', 'Total:', '.'),
    'it': ('Fattura:', 'Data:', 'Fornitore:', 'Valuta:', 'Totale:', ','),
}
EXAMPLES = {
    'example-en.pdf': ('en', [
        'FICTIONAL DEMONSTRATION - NOT A TAX DOCUMENT',
        'Invoice No: 0007', 'Date: 2026-10-20', 'Supplier: Example Studio',
        'Currency: EUR', 'ITEMS',
        'ITEM | Sample notebooks | 2 | 12.50 | 5.00 | 30.00',
        'ITEM | Sample envelopes | 1 | 10.00 | 2.00 | 12.00',
        'Total: 42.00',
    ]),
    'example-it.pdf': ('it', [
        'DIMOSTRAZIONE FITTIZIA - NON E UN DOCUMENTO FISCALE',
        'Fattura: 0008', 'Data: 21/10/2026', 'Fornitore: Laboratorio Esempio',
        'Valuta: EUR', 'RIGHE',
        'ITEM | Cartelle campione | 3 | 10,00 | 6,60 | 36,60',
        'ITEM | Scatola campione | 1 | 5,00 | 1,10 | 6,10',
        # Deliberate discrepancy: source total must be preserved and flagged.
        'Totale: 43,00',
    ]),
}


def amount(raw, separator):
    if not re.fullmatch(r'\d+' + re.escape(separator) + r'\d{2}', raw):
        raise ValueError('Amount outside agreed decimal format: ' + raw)
    return Decimal(raw.replace(separator, '.'))


def parse_text(text, layout):
    labels = LAYOUTS[layout]
    lines = text.splitlines()
    values = []
    for label in labels[:5]:
        matches = [line[len(label):].strip() for line in lines if line.startswith(label)]
        if len(matches) != 1 or not matches[0]:
            raise ValueError('Missing or repeated required field: ' + label)
        values.append(matches[0])
    invoice, date, supplier, currency, total_raw = values
    total = amount(total_raw, labels[5])
    items, issues = [], []
    for line_number, line in enumerate(lines, 1):
        if not line.startswith('ITEM |'):
            continue
        parts = [part.strip() for part in line.split('|')]
        if len(parts) != 6:
            raise ValueError('Unrecognized item row; do not silently skip')
        _, description, quantity, unit, tax, gross = parts
        if not re.fullmatch(r'\d+', quantity) or not description:
            raise ValueError('Unsupported quantity or empty description')
        unit_n, tax_n, gross_n = [amount(v, labels[5]) for v in (unit, tax, gross)]
        if Decimal(quantity) * unit_n + tax_n != gross_n:
            issues.append({'type': 'line_total_mismatch', 'text_line': line_number})
        items.append({'description': description, 'quantity_raw': quantity,
                      'unit_raw': unit, 'tax_raw': tax, 'gross_raw': gross,
                      'gross_numeric': str(gross_n), 'text_line': line_number})
    if not items:
        raise ValueError('No supported line items; review source manually')
    calculated = sum((Decimal(i['gross_numeric']) for i in items), Decimal(0))
    if calculated != total:
        issues.append({'type': 'invoice_total_mismatch',
                       'source_total': total_raw, 'sum_of_lines': str(calculated)})
    return dict(invoice=invoice, date_raw=date, supplier=supplier,
                currency=currency, total_raw=total_raw, items=items, issues=issues)


def parse_pdf(path, layout):
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) != 1:
            raise ValueError('This demo supports one-page invoices only')
        text = pdf.pages[0].extract_text() or ''
    result = parse_text(text, layout)
    result.update(source_file=path.name, source_page=1,
                  sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    return result


def write_workbook(records, path):
    wb = Workbook()
    summary = wb.active
    summary.title = 'Review'
    summary.append(['Fictional demonstration; no client or tax-document claims'])
    summary.append(['Source file', 'Page', 'Invoice ID', 'Status', 'Issue details', 'SHA256'])
    for record in records:
        title = re.sub(r'[\\/*?:\[\]]', '_', record['invoice'])[:31] or 'Invoice'
        base, suffix = title, 2
        while title.casefold() in {s.casefold() for s in wb.sheetnames}:
            title = base[:26] + '_' + str(suffix)
            suffix += 1
        sheet = wb.create_sheet(title)
        for field in ['source_file', 'source_page', 'invoice', 'date_raw', 'supplier', 'currency', 'total_raw']:
            sheet.append([field, record[field]])
        sheet.append([])
        sheet.append(['Description', 'Qty as printed', 'Unit as printed', 'Tax amount as printed', 'Line total as printed', 'Source text line'])
        for item in record['items']:
            sheet.append([item[k] for k in ['description', 'quantity_raw', 'unit_raw', 'tax_raw', 'gross_raw', 'text_line']])
        summary.append([record['source_file'], record['source_page'], record['invoice'],
                        'REVIEW REQUIRED' if record['issues'] else 'CHECKS PASSED',
                        json.dumps(record['issues'], ensure_ascii=False), record['sha256']])
        if record['issues']:
            sheet['B7'].fill = PatternFill('solid', fgColor='FFF2CC')
            for cell in summary[summary.max_row]:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        for row in sheet:
            for cell in row:
                # Treat all source strings literally, including formula-like input.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                    cell.number_format = '@'
                if cell.row == 1:
                    cell.font = Font(bold=True)
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            sheet.column_dimensions[col].width = 28
    wb.save(path)


def run():
    records = []
    for filename, (layout, lines) in EXAMPLES.items():
        path = ROOT / filename
        pdf = canvas.Canvas(str(path))
        pdf.setFont('Courier', 9)
        for n, line in enumerate(lines):
            pdf.drawString(35, 800 - n * 22, line)
        pdf.save()
        records.append(parse_pdf(path, layout))
    write_workbook(records, ROOT / 'invoice-review.xlsx')
    report = {'demonstration_only': True, 'layouts': ['en', 'it'],
              'invoices': records, 'external_requests': 0,
              'limitations': 'Two exact one-page fictional layouts; no OCR, general layout detection or Excel desktop compatibility claim.'}
    (ROOT / 'extraction-report.json').write_text(json.dumps(report, indent=2, ensure_ascii=False)+'\n')
    print(json.dumps({'pdfs': len(records), 'line_items': sum(len(r['items']) for r in records),
                      'review_flags': sum(len(r['issues']) for r in records), 'outputs_written': True}))


if __name__ == '__main__':
    run()
